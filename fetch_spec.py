#!/usr/bin/env python3
"""
Downloads a tracking spec straight from a Lark sheet link into the dashboard's
spec folder (~/Documents/tracking-spec by default) as an .xlsx.

Every row is written at its REAL Lark row number (from lark-cli's [row=N]),
so "row 17" on the dashboard / in reports is row 17 on Lark too. A running
server in folder mode picks the new file up by itself within ~2s.

Usage:
  python3 fetch_spec.py "https://technify.sg.larksuite.com/wiki/<token>?sheet=<id>"
  python3 fetch_spec.py "<link without ?sheet=>"      # uses the "Tracking..." tab
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl

from lark_sheet import first_sheet_id, parse_sheet, read_sheet, sheet_id_from_url


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("url", help="Link Lark (wiki hoặc sheets), có thể kèm ?sheet=<id>")
    ap.add_argument("--sheet-id", default=None, help="Sub-sheet id (mặc định: lấy từ ?sheet= trong link)")
    ap.add_argument("--out-dir", default="~/Documents/tracking-spec")
    args = ap.parse_args()

    sheet_id = args.sheet_id or sheet_id_from_url(args.url) or first_sheet_id(args.url)
    data = read_sheet(args.url, sheet_id)
    header_row, headers, _, rows = parse_sheet(data)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tracking Auto"
    for j, name in enumerate(headers, start=1):
        ws.cell(row=header_row, column=j, value=name)
    for r in rows:
        for j, name in enumerate(headers, start=1):
            val = r["cells"].get(name, "")
            if name and val != "":
                ws.cell(row=r["row_num"], column=j, value=val)

    out_dir = Path(args.out_dir).expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"lark-{sheet_id}.xlsx"
    wb.save(out)
    # Remember where this spec came from, so push_to_lark.py can write back
    # to the same sheet without the link being passed again.
    (out_dir / f"lark-{sheet_id}.source.json").write_text(
        json.dumps({"url": args.url, "sheet_id": sheet_id}, ensure_ascii=False), encoding="utf-8")

    events = sum(1 for r in rows if r["cells"].get("Event name", "").strip())
    print(json.dumps({"ok": True, "file": str(out), "sheet_id": sheet_id, "header_row": header_row,
                      "event_rows": events, "columns": [h for h in headers if h]},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    try:
        main()
    except RuntimeError as e:
        print(f"[error] {e}", file=sys.stderr)
        sys.exit(1)
