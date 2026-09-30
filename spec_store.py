import re
from pathlib import Path

import openpyxl

GLOBAL_ALLOWED_KEYS = {
    "screen_name", "prev_screen_name", "session_number",
    "overlay_name", "button_name", "timestamp",
    # Sent by the apps on every button_click but not declared in the specs;
    # QA decided (2026-09-29) not to count it as an unexpected param.
    "is_paywall",
}

_SPLIT_RE = re.compile(r',(?![^(]*\))')
_KEY_RE = re.compile(r'^\s*([A-Za-z0-9_\.]+)\s*[:=]')


def parse_expected_keys(bundle_params_text):
    if not bundle_params_text:
        return set()
    keys = set()
    for segment in _SPLIT_RE.split(bundle_params_text):
        m = _KEY_RE.match(segment)
        if m:
            keys.add(m.group(1).strip())
    return keys


def split_values(text):
    """A spec cell can hold several values (newline / comma / semicolon / slash
    separated), e.g. one screen_show row that covers ScrScanOptionPhoto + Video
    + Audio, or "ScrTranslate/ScrVoice/ScrCamera"."""
    return [v.strip() for v in re.split(r'[\n,;/]+', text or "") if v.strip()]


GENERIC_PARAMS = ("button_name", "screen_name", "overlay_name")


def custom_params(spec_row):
    """Param names that aren't the generic identifiers — e.g. "input_language".
    Two rows can share event + button and differ only by this param, so it
    has to be present in the bundle for the row to be a clean match."""
    return [p for p in split_values(spec_row.param_name) if p not in GENERIC_PARAMS]


class SpecRow:
    __slots__ = (
        "screen", "event_name", "param_name", "button_overlay_name",
        "screen_name", "overlay_name", "bundle_params_text",
        "expected_keys", "definition", "app_version", "row_num",
    )

    def __init__(self, **kw):
        for k, v in kw.items():
            setattr(self, k, v)

    def identifier_values(self):
        if self.event_name in ("screen_show", "screen_hide"):
            return split_values(self.screen_name)
        if self.event_name in ("overlay_show", "overlay_hide"):
            return split_values(self.overlay_name)
        return split_values(self.button_overlay_name) or split_values(self.param_name)

    def identifier_label(self):
        return ", ".join(self.identifier_values())


# Sheets in the wild use slightly different header names for the same concept
# (our own template vs. the team's live "Tracking Auto" Lark sheet). Each list
# is tried in order; the first header that actually exists in the sheet wins.
HEADER_ALIASES = {
    "screen": ["Feature", "Screen"],
    "event_name": ["Event name"],
    "param_name": ["Param name"],
    "button_overlay_name": ["Button name", "Button / Overlay name"],
    "screen_name": ["Screen name"],
    "overlay_name": ["Overlay name"],
    "bundle_params_text": ["Valid Value", "Bundle Params"],
    "definition": ["Định nghĩa", "Định nghĩa / Trigger"],
    "app_version": ["App Version"],
}


def _newest_xlsx_in(dir_path):
    """Picks the most recently modified .xlsx in a folder, ignoring Excel/Office
    lock files (~$foo.xlsx) and in-progress browser downloads (.crdownload/.part/.tmp)."""
    candidates = [
        p for p in dir_path.glob("*.xlsx")
        if not p.name.startswith("~$")
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


class SpecStore:
    """`location` can be a single .xlsx file, or a folder — in which case the
    most recently modified .xlsx inside it is used automatically. Folder mode
    exists so a fresh browser download (which the browser renames instead of
    overwriting, e.g. "Untitled spreadsheet (1).xlsx") just works without the
    user having to rename or manually replace anything."""

    def __init__(self, location):
        self.location = Path(location)
        self.path = None       # the actual resolved .xlsx currently in use
        self.rows = []
        self.error = None
        self.mtime = None
        self.load()

    def _resolve_path(self):
        if self.location.is_dir():
            return _newest_xlsx_in(self.location)
        return self.location

    def has_changed(self):
        resolved = self._resolve_path()
        if resolved is None:
            return False
        try:
            mtime = resolved.stat().st_mtime
        except FileNotFoundError:
            return False
        return resolved != self.path or mtime != self.mtime

    def _pick_worksheet(self, wb):
        # Prefer a sheet whose name suggests it's the live tracking/spec sheet
        # (e.g. exported from Lark's "Tracking Auto" tab); otherwise fall back
        # to the first sheet that isn't an obvious side-tab like "Legend".
        skip_names = {"legend", "sheet2"}
        for ws in wb.worksheets:
            if "tracking" in ws.title.lower():
                return ws
        for ws in wb.worksheets:
            if ws.title.lower() not in skip_names:
                return ws
        return wb.active

    def load(self):
        resolved = self._resolve_path()
        if resolved is None:
            self.error = (
                f"Không tìm thấy file .xlsx nào trong thư mục spec: {self.location}"
                if self.location.is_dir() else
                f"Không tìm thấy file spec: {self.location}"
            )
            self.rows = []
            return

        self.path = resolved
        try:
            self.mtime = self.path.stat().st_mtime
        except FileNotFoundError:
            pass

        try:
            wb = openpyxl.load_workbook(self.path, data_only=True)
        except Exception as e:
            self.error = f"Không mở được file spec: {e}"
            self.rows = []
            return

        ws = self._pick_worksheet(wb)

        header_row_idx = None
        headers = {}
        max_scan = min(ws.max_row, 10)
        for r in range(1, max_scan + 1):
            values = [c.value for c in ws[r]]
            if any(v and str(v).strip() == "Event name" for v in values):
                header_row_idx = r
                for idx, v in enumerate(values, start=1):
                    if v:
                        headers[str(v).strip()] = idx
                break

        if header_row_idx is None:
            self.error = "Không tìm thấy header 'Event name' trong file spec (kiểm tra lại sheet)."
            self.rows = []
            return

        def col_for(field):
            for alias in HEADER_ALIASES[field]:
                if alias in headers:
                    return headers[alias]
            return None

        def cell(r, field):
            c = col_for(field)
            if not c:
                return ""
            v = ws.cell(row=r, column=c).value
            return str(v).strip() if v is not None else ""

        rows = []
        last_feature = ""
        for r in range(header_row_idx + 1, ws.max_row + 1):
            event_name = cell(r, "event_name")
            if not event_name:
                continue
            bundle_text = cell(r, "bundle_params_text")
            # "Feature" is a merged cell in the real sheet: only the first row of a
            # group carries the value, the rest read back blank -> carry it down.
            feature = cell(r, "screen") or last_feature
            last_feature = feature
            rows.append(SpecRow(
                screen=feature,
                event_name=event_name,
                param_name=cell(r, "param_name"),
                button_overlay_name=cell(r, "button_overlay_name"),
                screen_name=cell(r, "screen_name"),
                overlay_name=cell(r, "overlay_name"),
                bundle_params_text=bundle_text,
                expected_keys=parse_expected_keys(bundle_text),
                definition=cell(r, "definition"),
                app_version=cell(r, "app_version"),
                row_num=r,
            ))
        self.rows = rows
        self.error = None

    def candidates(self, event_name):
        return [r for r in self.rows if r.event_name == event_name]


def concrete(values):
    """Drops placeholders like "(Screen tương ứng)" — real screen/overlay
    names never contain spaces or parentheses."""
    return [v for v in values if not re.search(r"[\s()]", v)]


_TYPE_WORDS = {"string", "int", "integer", "long", "number", "double", "float", "boolean", "bool"}


def valid_values(spec_row):
    """The closed list of allowed values for the row's custom param, taken
    from "Valid Value" (e.g. "1,2,3,4,5" or "EN, ES, VI"). None when the cell
    is open-ended ("EN, ES, PT, ...", "(rating content)"), key:value style,
    or when the row has no single custom param to check it against."""
    if len(custom_params(spec_row)) != 1 or spec_row.expected_keys:
        return None
    text = spec_row.bundle_params_text or ""
    if not text.strip() or "..." in text or "…" in text or "(" in text:
        return None
    tokens = [t for t in split_values(text) if t.lower() not in _TYPE_WORDS]
    if not tokens or any(re.search(r"\s", t) for t in tokens):
        return None
    return tokens


def score_match(spec_row, bundle):
    score = 0
    identifier_ok = False

    if spec_row.event_name in ("screen_show", "screen_hide"):
        if bundle.get("screen_name") in split_values(spec_row.screen_name):
            identifier_ok = True
            score += 3
    elif spec_row.event_name in ("overlay_show", "overlay_hide"):
        if bundle.get("overlay_name") in split_values(spec_row.overlay_name):
            identifier_ok = True
            score += 3
    else:
        targets = split_values(spec_row.button_overlay_name)
        params = custom_params(spec_row)
        custom_param = params[0] if params else None
        if targets:
            if bundle.get("button_name") in targets:
                identifier_ok = True
                score += 3
            elif custom_param and bundle.get(custom_param):
                # e.g. param_name="aiphoto_style" identifies the row by key
                # presence, since the logged value (a style code) won't equal
                # the human-readable button label in `targets`.
                identifier_ok = True
                score += 2
            elif any(t in bundle.values() for t in targets):
                identifier_ok = True
                score += 1
        elif custom_param and custom_param in bundle:
            identifier_ok = True
            score += 1

    if not identifier_ok:
        return None

    # The same button name is reused on many screens/overlays (BtnBack is on
    # half the app). A row that names where it lives only covers clicks there:
    # BtnBack logged on ScrAIVideoResult must not pass a row for ScrSetting.
    # A click logged with no overlay_name on a row that requires an overlay
    # stays a candidate (maybe the app forgot overlay_name), but as PARTIAL.
    context_missing = set()
    if spec_row.event_name not in ("screen_show", "screen_hide"):
        screens = concrete(split_values(spec_row.screen_name))
        if screens and bundle.get("screen_name"):
            if bundle["screen_name"] not in screens:
                return None
            score += 1
    if spec_row.event_name not in ("overlay_show", "overlay_hide"):
        overlays = concrete(split_values(spec_row.overlay_name))
        if overlays:
            if bundle.get("overlay_name"):
                if bundle["overlay_name"] not in overlays:
                    return None
                score += 2
            else:
                context_missing.add("overlay_name")

    bundle_keys = set(bundle.keys())
    params = set(custom_params(spec_row))
    ignore = GLOBAL_ALLOWED_KEYS | params
    missing = (spec_row.expected_keys | params | context_missing) - bundle_keys
    extra = bundle_keys - spec_row.expected_keys - ignore
    wrong = []
    allowed = valid_values(spec_row)
    if allowed:
        key = custom_params(spec_row)[0]
        val = bundle.get(key)
        if val is not None and str(val).strip().lower() not in {a.lower() for a in allowed}:
            wrong.append(f"{key}={val}")
    score -= len(missing) + len(extra) + len(wrong)
    return {"row": spec_row, "score": score, "missing": sorted(missing), "extra": sorted(extra),
            "wrong": wrong}


def match_event(spec_store, event_name, bundle):
    candidates = spec_store.candidates(event_name)
    scored = []
    for row in candidates:
        m = score_match(row, bundle)
        if m:
            scored.append(m)

    if not scored:
        return {
            "verdict": "UNKNOWN",
            "spec_row": None,
            "missing": [],
            "extra": [],
            "wrong": [],
            "also_rows": [],
        }

    best = max(scored, key=lambda m: m["score"])
    clean = lambda m: not m["missing"] and not m["extra"] and not m["wrong"]
    verdict = "MATCH" if clean(best) else "PARTIAL"
    # Sheets sometimes list the same event twice (e.g. once under onboarding,
    # once under settings). One firing that fits both equally well counts for both.
    also = [m["row"].row_num for m in scored
            if m is not best and m["score"] == best["score"] and clean(m) == clean(best)]
    return {
        "verdict": verdict,
        "spec_row": best["row"],
        "missing": best["missing"],
        "extra": best["extra"],
        "wrong": best["wrong"],
        "also_rows": also,
    }
