"""Session report builders (xlsx / md / csv) from the dashboard's in-memory state.

`data` shape (assembled by server.py):
  meta:     {generated_at, device, packages, spec_file, first_event, last_event, event_counts}
  coverage: {"rows": [...], "summary": {...}}   (same as /api/coverage)
  events:   [event payloads]
  errors:   [error payloads]
"""
import csv
import io
import json
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
STATUS_LABEL = {"PASSED": "Passed", "FAILED": "Failed", "IN_PROGRESS": "Đang test", "NOT_TESTED": "Chưa test",
                "BLOCKED": "Không test được"}
STATUS_FILL = {"PASSED": "C6EFCE", "FAILED": "FFC7CE", "IN_PROGRESS": "FFEB9C", "NOT_TESTED": "E7E6E6",
               "BLOCKED": "D9D2E9"}
STATUS_ORDER = ["FAILED", "IN_PROGRESS", "BLOCKED", "NOT_TESTED", "PASSED"]


def _pid_tid(ev):
    return f"{ev.get('pid')}-{ev.get('tid')}" if ev.get("pid") else ""


def _note(row):
    parts = []
    if row["reason"]:
        parts.append(row["reason"])
    if row.get("flaky_reason"):
        parts.append("Có lần sai, test lại thì đúng: " + row["flaky_reason"])
    if row.get("qa_note"):
        parts.append(row["qa_note"])
    if row["pending"] and row["status"] != "NOT_TESTED":
        parts.append(f"Chưa thấy ({len(row['pending'])}/{len(row['values'])}): " + ", ".join(row["pending"]))
    return " | ".join(parts)


def _seen(row):
    if not (row["match"] or row["partial"]):
        return ""
    return f"✓{row['match']}" + (f" ⚠{row['partial']}" if row["partial"] else "")


def unknown_events(events):
    """Events that fired but matched no spec row, grouped by (event, identifier)."""
    groups = OrderedDict()
    for ev in events:
        if ev.get("verdict") != "UNKNOWN":
            continue
        name = ev["event_name"]
        b = ev.get("bundle", {})
        ident = (b.get("screen_name") if name.startswith("screen_")
                 else b.get("overlay_name") if name.startswith("overlay_")
                 else b.get("button_name")) or ""
        g = groups.setdefault((name, ident), {"count": 0, "sample": b, "first": ev.get("log_time")})
        g["count"] += 1
    return [{"event_name": k[0], "identifier": k[1], **v} for k, v in groups.items()]


# ------------------------------------------------------------------ xlsx
def _style_header(ws, row=1):
    for c in ws[row]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _finish_sheet(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.font.name != FONT:
                c.font = Font(name=FONT, size=10, bold=c.font.bold, color=c.font.color)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    _style_header(ws)


def build_xlsx(data):
    meta, cov = data["meta"], data["coverage"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Tổng quan"
    sm = cov["summary"]
    tested = sm["PASSED"] + sm["FAILED"] + sm["IN_PROGRESS"]
    rows = [
        ("Thời điểm xuất", meta["generated_at"]),
        ("Thiết bị", meta["device"]),
        ("App (package)", ", ".join(meta["packages"]) or "—"),
        ("File spec", meta["spec_file"]),
        ("Phiên test", f"{meta['first_event'] or '—'}  →  {meta['last_event'] or '—'}"),
        ("", ""),
        ("Tổng dòng spec", sm["total"]),
        ("Passed", sm["PASSED"]),
        ("Failed", sm["FAILED"]),
        ("Đang test (mới thấy một phần)", sm["IN_PROGRESS"]),
        ("Không test được", sm.get("BLOCKED", 0)),
        ("Chưa test", sm["NOT_TESTED"]),
        ("Đã đụng tới", f"{tested}/{sm['total']}" + (f" ({round(tested * 100 / sm['total'])}%)" if sm["total"] else "")),
        ("", ""),
        ("Tổng event bắn ra", sum(meta["event_counts"].values())),
        ("  MATCH", meta["event_counts"].get("MATCH", 0)),
        ("  PARTIAL", meta["event_counts"].get("PARTIAL", 0)),
        ("  UNKNOWN (không có trong spec)", meta["event_counts"].get("UNKNOWN", 0)),
        ("Lỗi/crash của app (Errors)", len(data["errors"])),
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows():
        row[0].font = Font(name=FONT, bold=True, size=10)
        row[1].font = Font(name=FONT, size=10)
        row[1].alignment = Alignment(horizontal="left", wrap_text=True)
    fills = {"Passed": "PASSED", "Failed": "FAILED", "Đang test (mới thấy một phần)": "IN_PROGRESS", "Chưa test": "NOT_TESTED",
             "Không test được": "BLOCKED"}
    for row in ws.iter_rows():
        if row[0].value in fills:
            row[1].fill = PatternFill("solid", fgColor=STATUS_FILL[fills[row[0].value]])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80

    wc = wb.create_sheet("Coverage")
    wc.append(["Trạng thái", "Feature", "Event", "Định danh", "Định nghĩa", "Đã thấy", "Ghi chú", "Dòng trong spec"])
    ordered = sorted(cov["rows"], key=lambda r: (STATUS_ORDER.index(r["status"]), r["row_num"]))
    for r in ordered:
        wc.append([STATUS_LABEL[r["status"]], r["feature"], r["event_name"], r["identifier"],
                   r["definition"], _seen(r), _note(r), r["row_num"]])
        wc.cell(row=wc.max_row, column=1).fill = PatternFill("solid", fgColor=STATUS_FILL[r["status"]])
    _finish_sheet(wc, [12, 16, 16, 34, 50, 10, 50, 10])

    wu = wb.create_sheet("Event không có trong spec")
    wu.append(["Event", "Định danh", "Số lần", "Lần đầu", "Bundle mẫu"])
    for g in unknown_events(data["events"]):
        wu.append([g["event_name"], g["identifier"], g["count"], g["first"] or "",
                   json.dumps(g["sample"], ensure_ascii=False)])
    _finish_sheet(wu, [20, 30, 8, 24, 80])

    we = wb.create_sheet("Events")
    we.append(["Thời gian", "PID-TID", "Package", "Event", "Định danh", "Verdict", "Bundle", "Ghi chú"])
    for ev in data["events"]:
        b = ev.get("bundle", {})
        ident = (b.get("screen_name") if ev["event_name"].startswith("screen_")
                 else b.get("overlay_name") if ev["event_name"].startswith("overlay_")
                 else b.get("button_name")) or ""
        note = " | ".join(x for x in (
            ("Thiếu: " + ", ".join(ev["missing"])) if ev.get("missing") else "",
            ("Lạ: " + ", ".join(ev["extra"])) if ev.get("extra") else "") if x)
        we.append([ev.get("log_time") or "", _pid_tid(ev), ev.get("package", ""), ev["event_name"], ident,
                   ev["verdict"], json.dumps(b, ensure_ascii=False), note])
    _finish_sheet(we, [24, 14, 30, 20, 26, 10, 80, 30])

    wr = wb.create_sheet("Errors")
    wr.append(["Thời gian", "PID-TID", "Level", "Tag", "Package", "Nội dung"])
    for er in data["errors"]:
        wr.append([er.get("log_time") or "", _pid_tid(er), "FATAL" if er["level"] == "F" else "ERROR",
                   er["tag"], er["package"], er["message"]])
    _finish_sheet(wr, [24, 14, 8, 22, 30, 100])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ csv
def build_csv(data):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Trạng thái", "Feature", "Event", "Định danh", "Định nghĩa", "Đã thấy", "Ghi chú", "Dòng trong spec"])
    ordered = sorted(data["coverage"]["rows"], key=lambda r: (STATUS_ORDER.index(r["status"]), r["row_num"]))
    for r in ordered:
        w.writerow([STATUS_LABEL[r["status"]], r["feature"], r["event_name"], r["identifier"],
                    r["definition"], _seen(r), _note(r), r["row_num"]])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8


# ------------------------------------------------------------------ markdown
def _md(text):
    return str(text).replace("|", "\\|").replace("\n", " ")


def build_md(data):
    meta, cov = data["meta"], data["coverage"]
    sm = cov["summary"]
    tested = sm["PASSED"] + sm["FAILED"] + sm["IN_PROGRESS"]
    by = {k: [r for r in cov["rows"] if r["status"] == k] for k in STATUS_ORDER}
    out = [
        "# Báo cáo test tracking",
        "",
        f"- **Thời điểm xuất:** {meta['generated_at']}",
        f"- **Thiết bị:** {meta['device']}",
        f"- **App:** {', '.join(meta['packages']) or '—'}",
        f"- **Spec:** {meta['spec_file']}",
        f"- **Phiên test:** {meta['first_event'] or '—'} → {meta['last_event'] or '—'}",
        "",
        "## Tổng quan",
        "",
        "| Passed | Failed | Đang test | Không test được | Chưa test | Tổng |",
        "|---|---|---|---|---|---|",
        f"| {sm['PASSED']} | {sm['FAILED']} | {sm['IN_PROGRESS']} | {sm.get('BLOCKED', 0)} | {sm['NOT_TESTED']} | {sm['total']} |",
        "",
        f"Đã đụng tới **{tested}/{sm['total']}** dòng spec"
        + (f" ({round(tested * 100 / sm['total'])}%)." if sm["total"] else "."),
        f"Event bắn ra: {sum(meta['event_counts'].values())} "
        f"(MATCH {meta['event_counts'].get('MATCH', 0)}, PARTIAL {meta['event_counts'].get('PARTIAL', 0)}, "
        f"UNKNOWN {meta['event_counts'].get('UNKNOWN', 0)}). Lỗi/crash của app: {len(data['errors'])}.",
        "",
    ]

    def table(title, rows, cols):
        if not rows:
            return
        out.append(f"## {title} ({len(rows)})")
        out.append("")
        out.append("| " + " | ".join(h for h, _ in cols) + " |")
        out.append("|" + "---|" * len(cols))
        for r in rows:
            out.append("| " + " | ".join(_md(f(r)) for _, f in cols) + " |")
        out.append("")

    base = [("Feature", lambda r: r["feature"]), ("Event", lambda r: r["event_name"]),
            ("Định danh", lambda r: r["identifier"])]
    table("Failed", by["FAILED"], base + [("Ghi chú", _note)])
    table("Đang test", by["IN_PROGRESS"], base + [("Ghi chú", _note)])
    table("Không test được", by["BLOCKED"], base + [("Lý do", _note)])
    table("Chưa test", by["NOT_TESTED"], base + [("Định nghĩa", lambda r: r["definition"])])
    table("Passed", by["PASSED"], base + [("Đã thấy", _seen)])

    unk = unknown_events(data["events"])
    table("Event bắn ra nhưng không có trong spec", unk,
          [("Event", lambda g: g["event_name"]), ("Định danh", lambda g: g["identifier"]),
           ("Số lần", lambda g: g["count"]),
           ("Bundle mẫu", lambda g: json.dumps(g["sample"], ensure_ascii=False))])
    return "\n".join(out).encode("utf-8")
